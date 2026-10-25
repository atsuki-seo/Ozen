# SPDX-License-Identifier: MIT OR Apache-2.0

import sys

sys.dont_write_bytecode = True

import glob
import os
import shutil
import subprocess
import tempfile
import unittest

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
EXTRACT = os.path.join(ROOT, "tools", "fct", "extract.py")
REAL = sorted(glob.glob(os.path.join(ROOT, "data", "mext-fct-2023", "source", "*_02.xlsx")))

from tools.fct.extract import HEADERS, read_refuse  # noqa: E402


def make_xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表全体"
    for cell, text in HEADERS.items():
        ws[cell] = text
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            ws.cell(row=13 + i, column=1 + j, value=v)
    wb.save(path)


class Extract(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="ozen-fct-")

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_reads_fixture_xlsx(self):
        path = os.path.join(self.tmp, "t.xlsx")
        make_xlsx(path, [("01", "00001", "0001", "食品A", 5), ("01", "00002", "0002", "食品B", None),
                         ("01", "00003", "0003", "食品C", "Tr"), ("01", "00004", "0004", "食品D", "(10)")])
        got = read_refuse(path)
        self.assertEqual(got["00001"], ("食品A", "5"))
        self.assertEqual(got["00002"], ("食品B", ""))
        self.assertEqual(got["00003"], ("食品C", "Tr"))
        self.assertEqual(got["00004"], ("食品D", "(10)"))

    def test_rejects_wrong_headers(self):
        path = os.path.join(self.tmp, "t.xlsx")
        make_xlsx(path, [("01", "00001", "0001", "食品A", 5)])
        wb = openpyxl.load_workbook(path)
        wb["表全体"]["E12"] = "WATER"
        wb.save(path)
        with self.assertRaises(ValueError):
            read_refuse(path)

    def test_writes_sorted_csv_and_checks(self):
        xlsx = os.path.join(self.tmp, "t.xlsx")
        make_xlsx(xlsx, [("01", "00009", "0009", "食品Z", 1), ("01", "00001", "0001", "食品A", 5),
                         ("01", "00005", "0005", "食品E", 3)])
        vocab = os.path.join(self.tmp, "vocab")
        os.makedirs(vocab)
        with open(os.path.join(vocab, "ingredients.yaml"), "w", encoding="utf-8") as f:
            f.write('ingredients:\n  B: {props: {refuse: {source: mext-fct-2023, food: "00009", field: refuse}}}\n'
                    '  A: {props: {refuse: {source: mext-fct-2023, food: "00001", field: refuse}}}\n')
        out = os.path.join(self.tmp, "refuse.csv")
        readme = os.path.join(self.tmp, "README.md")
        r = subprocess.run([sys.executable, "-B", EXTRACT, "--xlsx", xlsx, "--vocab", vocab, "--out", out],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        with open(out, encoding="utf-8") as f:
            self.assertEqual(f.read(), "food,name,refuse\n00001,食品A,5\n00009,食品Z,1\n")
        import hashlib
        with open(xlsx, "rb") as f:
            digest = hashlib.sha256(f.read()).hexdigest()
        with open(readme, "w", encoding="utf-8") as f:
            f.write(f"- SHA-256: {digest}\n")
        args = [sys.executable, "-B", EXTRACT, "--check", "--xlsx", xlsx, "--vocab", vocab, "--out", out, "--readme", readme]
        r = subprocess.run(args, capture_output=True, text=True)
        self.assertEqual(r.stdout.strip().splitlines()[-1], "fct: 2 rows, 0 mismatches; sha256 ok; order ok; refs ok; headers ok")
        with open(out, "w", encoding="utf-8") as f:
            f.write("food,name,refuse\n00001,食品A,6\n00009,食品Z,1\n")
        r = subprocess.run(args, capture_output=True, text=True)
        self.assertEqual(r.returncode, 1)
        self.assertIn("fct: 2 rows, 1 mismatches", r.stdout)

    @unittest.skipUnless(REAL, "the MEXT Excel file is not in data/mext-fct-2023/source/")
    def test_matches_real_xlsx(self):
        got = read_refuse(REAL[0])
        self.assertEqual(got["06153"][1], "6")
        self.assertIn("たまねぎ", got["06153"][0])
        self.assertEqual(got["07107"][1], "40")
        self.assertIn("バナナ", got["07107"][0])
        self.assertEqual(got["06061"][1], "15")
        self.assertIn("キャベツ", got["06061"][0])


if __name__ == "__main__":
    unittest.main()
