#!/usr/bin/env python3
# SPDX-License-Identifier: MIT OR Apache-2.0

import sys

sys.dont_write_bytecode = True

import argparse
import csv
import glob
import hashlib
import io
import os
import re

import openpyxl
import yaml

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SHEET = "表全体"
HEADERS = {"B2": "食　品　番　号", "D3": "食　品　名", "E3": "廃　棄　率", "E12": "REFUSE"}
FIRST_ROW = 13


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_refuse(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[SHEET]
        for ref, want in HEADERS.items():
            got = ws[ref].value
            if got != want:
                raise ValueError(f"header {ref} is {got!r}, expected {want!r}")
        out = {}
        for row in ws.iter_rows(min_row=FIRST_ROW, max_col=5, values_only=True):
            food = cell_text(row[1])
            if re.fullmatch(r"[0-9]{5}", food):
                out[food] = (cell_text(row[3]), cell_text(row[4]))
        return out
    finally:
        wb.close()


def table_refs(vocab_dir):
    refs = set()

    def walk(x):
        if isinstance(x, dict):
            if x.get("source") == "mext-fct-2023" and "food" in x:
                refs.add(str(x["food"]))
            for v in x.values():
                walk(v)
        elif isinstance(x, list):
            for v in x:
                walk(v)
    for path in glob.glob(os.path.join(vocab_dir, "*.yaml")):
        with open(path, encoding="utf-8") as f:
            walk(yaml.safe_load(f))
    return refs


def render(rows, refs):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["food", "name", "refuse"])
    for food in sorted(refs):
        name, refuse = rows.get(food, ("", ""))
        w.writerow([food, name, refuse])
    return buf.getvalue()


def sha256(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def check(args):
    try:
        rows = read_refuse(args.xlsx)
        headers = "ok"
    except (ValueError, KeyError) as e:
        print(e, file=sys.stderr)
        rows, headers = {}, "failed"
    refs = table_refs(args.vocab)
    with open(args.out, encoding="utf-8") as f:
        current = list(csv.DictReader(f))
    mismatches = 0
    for r in current:
        want = rows.get(r["food"])
        if want is None or (r["name"], r["refuse"]) != want:
            mismatches += 1
            print(f"mismatch: {r['food']} has {r['name']!r},{r['refuse']!r}, Excel has {want!r}", file=sys.stderr)
    foods = [r["food"] for r in current]
    order = "ok" if foods == sorted(foods) and len(set(foods)) == len(foods) else "failed"
    ref_state = "ok" if set(foods) == refs and refs <= set(rows) else "failed"
    with open(args.readme, encoding="utf-8") as f:
        m = re.search(r"^- SHA-256: ([0-9a-f]{64})$", f.read(), re.M)
    sha = "ok" if m and m.group(1) == sha256(args.xlsx) else "failed"
    print(f"fct: {len(current)} rows, {mismatches} mismatches; sha256 {sha}; order {order}; refs {ref_state}; headers {headers}")
    return 0 if not mismatches and sha == order == ref_state == headers == "ok" else 1


def main(argv=None):
    data = os.path.join(ROOT, "data", "mext-fct-2023")
    ap = argparse.ArgumentParser(prog="extract.py")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--vocab", default=os.path.join(ROOT, "spec", "vocab"))
    ap.add_argument("--out", default=os.path.join(data, "refuse.csv"))
    ap.add_argument("--readme", default=os.path.join(data, "README.md"))
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args(argv)
    if args.check:
        return check(args)
    rows = read_refuse(args.xlsx)
    refs = table_refs(args.vocab)
    unknown = sorted(refs - set(rows))
    if unknown:
        print("food numbers not in the Excel file: " + ", ".join(unknown), file=sys.stderr)
        return 1
    with open(args.out, "w", encoding="utf-8", newline="") as f:
        f.write(render(rows, refs))
    print(f"fct: wrote {len(refs)} rows to {os.path.relpath(args.out, ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
